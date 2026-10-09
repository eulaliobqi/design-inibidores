"""Pos-processamento comum dos .docx (margens, numeracao de linhas/paginas, fontes, tabelas).
Copiado de build_docx_pt.py (etapa 6) para o build em ingles; o PT segue com sua copia."""
from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


def postprocess(out, FIRST_TITLE):
    doc = Document(str(out))
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(2.5)
        s.top_margin = s.bottom_margin = Cm(2.5)
        ln = OxmlElement("w:lnNumType")
        ln.set(qn("w:countBy"), "1")
        ln.set(qn("w:restart"), "continuous")
        s._sectPr.append(ln)
        # numero de pagina no rodape
        p = s.footer.paragraphs[0] if s.footer.paragraphs else s.footer.add_paragraph()
        p.alignment = 1
        for tag, txt in (("begin", None), (None, "PAGE"), ("end", None)):
            r = p.add_run()
            if tag:
                fc = OxmlElement("w:fldChar")
                fc.set(qn("w:fldCharType"), tag)
                r._r.append(fc)
            else:
                it = OxmlElement("w:instrText")
                it.set(qn("xml:space"), "preserve")
                it.text = txt
                r._r.append(it)

    for name in ("Normal", "Body Text", "First Paragraph", "Compact"):
        if name in [s.name for s in doc.styles]:
            st = doc.styles[name]
            st.font.name = "Times New Roman"
            st.font.size = Pt(12)
            st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
            if name in ("Body Text", "First Paragraph"):
                st.paragraph_format.line_spacing = 1.5
                st.paragraph_format.space_after = Pt(6)
    for name, size in (("Heading 1", 16), ("Heading 2", 14), ("Heading 3", 12), ("Title", 18)):
        if name in [s.name for s in doc.styles]:
            st = doc.styles[name]
            st.font.name = "Times New Roman"
            st.font.size = Pt(size)
            st.font.bold = True
            st.font.color.rgb = None
            st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    if "Pendente" in [s.name for s in doc.styles]:
        doc.styles["Pendente"].font.highlight_color = WD_COLOR_INDEX.YELLOW
        doc.styles["Pendente"].font.bold = True

    for t in doc.tables:
        t.style = "Table Grid" if "Table Grid" in [s.name for s in doc.styles] else t.style
        tbl = t._tbl
        tblPr = tbl.tblPr
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            e = OxmlElement(f"w:{edge}")
            e.set(qn("w:val"), "single")
            e.set(qn("w:sz"), "4")
            e.set(qn("w:color"), "808080")
            borders.append(e)
        tblPr.append(borders)
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.line_spacing = 1.0
                    for r in p.runs:
                        r.font.size = Pt(9)
        for cell in t.rows[0].cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    r.bold = True
    for p in doc.paragraphs:
        if p.text.startswith(FIRST_TITLE):
            p.paragraph_format.page_break_before = True
            break

    doc.save(str(out))

"""Ajustes para o SIA sobre o arquivo EDITADO pelo autor (preserva as logos que ele pos na capa).

- capa: autores, e-mails, setor, categoria (as logos SIA e UFV ficam onde estao)
- novo slide 2: ODS 2 (icone oficial + meta 2.4 + como o trabalho se encaixa)
- slide de referencias: 5 principais
- ultimo slide: financiamento e agradecimentos, com logotipos
- renumera os slides

Uso: python apresentacao_congresso/edit_sia.py   (fecha o PowerPoint antes)
"""
import copy
import shutil
from io import BytesIO
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

AQUI = Path(__file__).resolve().parent
ARQ = AQUI / "Apresentacao_congresso_v2.pptx"
shutil.copy2(ARQ, AQUI / "Apresentacao_congresso_v2_antes_SIA.pptx")

LARG, ALT = Inches(13.333), Inches(7.5)
MARGEM = Inches(0.78)
GRAF = RGBColor(0x1F, 0x23, 0x28)
AZUL = RGBColor(0x00, 0x72, 0xB2)
VERDE = RGBColor(0x2A, 0x9D, 0x8F)
LARA = RGBColor(0xD5, 0x5E, 0x00)
AMAR = RGBColor(0xA8, 0x70, 0x00)
CINZA = RGBColor(0x5A, 0x5F, 0x66)
CLARO = RGBColor(0x8A, 0x90, 0x97)
FONTE = "Segoe UI"


def caixa(slide, x, y, cx, cy, texto, tam, cor=GRAF, negrito=False, alinha=PP_ALIGN.LEFT,
          italico=False, espaco=1.0):
    tb = slide.shapes.add_textbox(x, y, cx, cy)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, linha in enumerate(texto.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = alinha
        p.line_spacing = espaco
        r = p.add_run()
        r.text = linha
        r.font.size = Pt(tam)
        r.font.bold = negrito
        r.font.italic = italico
        r.font.name = FONTE
        r.font.color.rgb = cor
    return tb


def regua(slide, x, y, cx, cor=AZUL, espessura=Pt(3)):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, cx, espessura)
    s.fill.solid()
    s.fill.fore_color.rgb = cor
    s.line.fill.background()
    s.shadow.inherit = False
    return s


def notas(slide, texto):
    tf = slide.notes_slide.notes_text_frame
    if tf is not None:
        tf.text = texto


def nova(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    return slide


def titulo(slide, frase):
    caixa(slide, MARGEM, Inches(0.52), LARG - 2 * MARGEM, Inches(1.0), frase, 25, GRAF, negrito=True, espaco=1.08)
    regua(slide, MARGEM, Inches(0.40), Inches(0.9))


def rodape(slide, texto):
    caixa(slide, MARGEM, ALT - Inches(0.55), LARG - 2 * MARGEM - Inches(0.7), Inches(0.4), texto, 10, CLARO, espaco=1.1)


def imagem(slide, arquivo, x, y, altura=None, largura=None):
    w, h = Image.open(arquivo).size
    if altura:
        cy = altura
        cx = int(altura * w / h)
    else:
        cx = largura
        cy = int(largura * h / w)
    return slide.shapes.add_picture(str(arquivo), x, y, cx, cy)


def chip(slide, x, y, cx, rotulo, texto, cor, alt=Inches(1.0), tam_t=15):
    regua(slide, x, y, Inches(0.06), cor, alt)
    caixa(slide, x + Inches(0.22), y - Inches(0.02), cx - Inches(0.22), Inches(0.4), rotulo, 20, cor, negrito=True)
    caixa(slide, x + Inches(0.22), y + Inches(0.4), cx - Inches(0.22), alt, texto, tam_t, GRAF, espaco=1.18)


prs = Presentation(ARQ)

# ============================================================== capa ====
s1 = prs.slides[0]
alvo = [sh for sh in s1.shapes if sh.has_text_frame and sh.text_frame.text.startswith("Eulálio Gutemberg")][0]
alvo.left, alvo.top, alvo.width, alvo.height = MARGEM, Inches(4.72), Inches(8.6), Inches(1.9)
tf = alvo.text_frame
tf.word_wrap = True
for p in list(tf.paragraphs)[1:]:
    p._p.getparent().remove(p._p)
linhas = [
    ("Eulálio Gutemberg Dias dos Santos (1º autor) · eulalio.santos@ufv.br", 14, True, GRAF),
    ("Orientadora: Maria Goreti de Almeida Oliveira · malmeida@ufv.br", 12.5, False, GRAF),
    ("Co-orientador: Otávio J. B. Brustolini · tavinbio@lncc.br", 12.5, False, GRAF),
    ("Colaborador: Daniel G. S. Paulo · daniel.paulo@ufv.br", 12.5, False, GRAF),
    ("Departamento de Bioquímica e Biologia Molecular (DBB) · Universidade Federal de Viçosa", 12.5, False, CINZA),
    ("Categoria: Pesquisa (pós-graduação)", 13, True, AZUL),
]
p0 = tf.paragraphs[0]
for r in list(p0.runs)[1:]:
    r._r.getparent().remove(r._r)
for i, (txt, tam, neg, cor) in enumerate(linhas):
    p = p0 if i == 0 else tf.add_paragraph()
    p.line_spacing = 1.18
    if i == 0:
        r = p.runs[0] if p.runs else p.add_run()
    else:
        r = p.add_run()
    r.text = txt
    r.font.size = Pt(tam)
    r.font.bold = neg
    r.font.name = FONTE
    r.font.color.rgb = cor

# ======================================================== novo slide: ODS ====
s = nova(prs)
titulo(s, "Este trabalho contribui para o ODS 2: fome zero e agricultura sustentável")
ods = AQUI / "imagens" / "ods" / "E_SDG_Icons-02.jpg"
imagem(s, ods, MARGEM, Inches(2.0), altura=Inches(3.6))
caixa(s, MARGEM, Inches(5.75), Inches(3.6), Inches(0.6),
      "ODS 2 · Fome zero e agricultura sustentável", 13, CINZA, alinha=PP_ALIGN.CENTER, espaco=1.15)
_x = MARGEM + Inches(4.2)
_cx = LARG - _x - MARGEM
regua(s, _x, Inches(2.0), Inches(0.06), AMAR, Inches(1.55))
caixa(s, _x + Inches(0.22), Inches(1.98), _cx - Inches(0.22), Inches(0.4), "Meta 2.4 (ONU)", 20, AMAR, negrito=True)
caixa(s, _x + Inches(0.22), Inches(2.42), _cx - Inches(0.22), Inches(1.2),
      "“Até 2030, garantir sistemas sustentáveis de produção de alimentos e implementar práticas agrícolas "
      "resilientes, que aumentem a produtividade e a produção […]”", 15, GRAF, italico=True, espaco=1.2)
caixa(s, _x, Inches(3.95), _cx, Inches(0.4), "Como o trabalho se encaixa", 18, GRAF, negrito=True)
itens = [
    ("Proteger a produtividade da soja", "contra lagartas-praga: pragas e patógenos tiram 21,4% da soja mundial", AZUL),
    ("Uma via biorracional", "peptídeos que atacam a digestão da lagarta, com potencial de reduzir a dependência de inseticidas", VERDE),
    ("Triagem computacional", "prioriza poucos candidatos antes da síntese e do ensaio, poupando tempo e reagentes", LARA),
]
for i, (t, d, cor) in enumerate(itens):
    y = Inches(4.45) + i * Inches(0.78)
    regua(s, _x, y, Inches(0.06), cor, Inches(0.62))
    caixa(s, _x + Inches(0.22), y - Inches(0.03), _cx - Inches(0.22), Inches(0.3), t, 15, cor, negrito=True)
    caixa(s, _x + Inches(0.22), y + Inches(0.27), _cx - Inches(0.22), Inches(0.4), d, 13, GRAF)
rodape(s, "ONU Brasil, ODS 2 e meta 2.4 (brasil.un.org/pt-br/sdgs/2)  ·  ícone oficial: sdgs.un.org  ·  "
          "Savary et al. 2019, Nat. Ecol. Evol. 3:430")
notas(s, (
    "20 s. Área temática do trabalho: ODS 2. Meta 2.4: sistemas sustentáveis de produção de alimentos e práticas "
    "agrícolas resilientes. Encaixe: proteger a soja de lagartas por via biorracional, com triagem computacional antes do "
    "ensaio. O potencial de reduzir inseticidas é uma expectativa, ainda não testada."))
# move para a posição 2
lst = prs.slides._sldIdLst
ids = list(lst)
no = ids[-1]
lst.remove(no)
lst.insert(1, no)

# ============================================ referencias: 5 principais ====
ref_slide = [sl for sl in prs.slides if any(sh.has_text_frame and sh.text_frame.text.strip() == "Referências" for sh in sl.shapes)][0]
import re as _re
caixas = [sh for sh in ref_slide.shapes if sh.has_text_frame and _re.match(r"^\d+\. ", sh.text_frame.text)]
caixas.sort(key=lambda z: z.left)
for sh in ref_slide.shapes:
    if sh.has_text_frame and sh.text_frame.text.strip() == "Referências":
        sh.text_frame.paragraphs[0].runs[0].text = "Principais referências"
cinco = [
    "Savary S et al. (2019) The global burden of pathogens and pests on major food crops. Nat Ecol Evol 3:430–439. doi:10.1038/s41559-018-0793-y",
    "Almeida Barros R et al. (2021) Small peptides inhibit gut trypsin-like proteases and impair Anticarsia gemmatalis survival and development. Pest Manag Sci 77:1714–1723. doi:10.1002/ps.6191",
    "Júnior NRS et al. (2026) Integrated transcriptomic and physiological analyses reveal disruption of digestive homeostasis induced by protease inhibitors in A. gemmatalis. Arch Insect Biochem Physiol 123:e70223. doi:10.1002/arch.70223",
    "Watson JL et al. (2023) De novo design of protein structure and function with RFdiffusion. Nature 620:1089–1100. doi:10.1038/s41586-023-06415-8",
    "Perona JJ, Craik CS (1995) Structural basis of substrate specificity in the serine proteases. Protein Sci 4:337–360. doi:10.1002/pro.5560040301",
]
c0 = caixas[0]
c0.left, c0.top, c0.width, c0.height = MARGEM, Inches(1.7), LARG - 2 * MARGEM, Inches(4.6)
tf = c0.text_frame
for p in list(tf.paragraphs)[1:]:
    p._p.getparent().remove(p._p)
p0 = tf.paragraphs[0]
for r in list(p0.runs)[1:]:
    r._r.getparent().remove(r._r)
for i, ref in enumerate(cinco):
    p = p0 if i == 0 else tf.add_paragraph()
    p.line_spacing = 1.2
    p.space_after = Pt(9)
    r = p.runs[0] if (i == 0 and p.runs) else p.add_run()
    r.text = f"{i + 1}. {ref}"
    r.font.size = Pt(15)
    r.font.name = FONTE
    r.font.color.rgb = GRAF
for sh in caixas[1:]:
    sh._element.getparent().remove(sh._element)
for sh in ref_slide.shapes:
    if sh.has_text_frame and sh.text_frame.text.startswith("Dados de resultados"):
        sh.text_frame.paragraphs[0].runs[0].text = (
            "Referências completas de cada slide nas anotações  ·  resultados: manuscrito do autor (em preparação)")
notas(ref_slide, (
    "Lista completa (22): Savary 2019; Oliveira 2014 (10.1016/j.cropro.2013.10.022); Carpane 2022 (10.1371/journal.pone.0271084); "
    "Ongaratto 2021 (10.1093/jee/toab197); Conab set/2026; Luckett 1999 (10.1006/jmbi.1999.2891); Werner & Wemmer 1992 "
    "(10.1021/bi00119a008); Song & Suh 1998 (10.1006/jmbi.1997.1469); Almeida Barros 2022 (10.1016/j.pestbp.2022.105188); "
    "Laskowski & Kato 1980 (10.1146/annurev.bi.49.070180.003113); Perona & Craik 1995; Hedstrom 2002 (10.1021/cr000033x); "
    "Almeida Barros 2021; Júnior 2026; Mariano 2026 (10.1002/arch.70224); Dow 1992 (10.1242/jeb.172.1.355); Watson 2023; "
    "Dauparas 2022 (10.1126/science.add2187); Passaro 2025 (10.1101/2025.06.14.659707); Abraham 2015 (10.1016/j.softx.2015.06.001); "
    "Huang & MacKerell 2013 (10.1002/jcc.23354); Vangone & Bonvin 2015 (10.7554/eLife.07454)."))

# ============================== ultimo slide: financiamento e agradecimentos ====
s = nova(prs)
titulo(s, "Financiamento e agradecimentos")
LG = AQUI / "imagens" / "logos"
caixa(s, MARGEM, Inches(1.6), Inches(6), Inches(0.35), "Apoio financeiro", 16, CINZA, negrito=True)
_h = Inches(1.05)
x = MARGEM
for nome, leg in (("fapemig.png", "FAPEMIG · bolsa 2025–2026"), ("capes.png", "CAPES"), ("cnpq.png", "CNPq")):
    pic = imagem(s, LG / nome, x, Inches(2.05), altura=_h)
    caixa(s, x, Inches(2.05) + _h + Inches(0.08), max(pic.width, Inches(2.0)), Inches(0.3), leg, 11, CLARO)
    x += max(pic.width, Inches(2.0)) + Inches(0.6)
regua(s, MARGEM, Inches(3.8), LARG - 2 * MARGEM, RGBColor(0xE3, 0xE6, 0xEA), Pt(1.5))
caixa(s, MARGEM, Inches(4.0), Inches(6), Inches(0.35), "Agradecimentos", 16, CINZA, negrito=True)
caixa(s, MARGEM, Inches(4.45), Inches(6.4), Inches(2.0),
      "À orientadora Maria Goreti de Almeida Oliveira, ao co-orientador Otávio J. B. Brustolini e ao "
      "colaborador Daniel G. S. Paulo; ao Departamento de Bioquímica e Biologia Molecular (DBB), ao BIOAGRO e à "
      "Universidade Federal de Viçosa; à FAPEMIG, à CAPES e ao CNPq pelo apoio.", 15, GRAF, espaco=1.3)
lab = imagem(s, Path(r"C:\Users\eulal\Desktop\logo-lab.png"), LARG - MARGEM - Inches(5.2), Inches(4.2), largura=Inches(5.2))
# UFV: reaproveita a imagem que o autor colocou na capa
imagem(s, AQUI / "imagens" / "logos" / "ufv_capa.png", LARG - MARGEM - Inches(2.0), Inches(1.95), altura=Inches(1.3))
rodape(s, "Apoio financeiro conforme o resumo submetido ao SIA: CAPES, CNPq e FAPEMIG (bolsa).")
notas(s, "10 s. Agradecer e abrir para perguntas.")

# ======================================================= renumera os slides ====
for i, sl in enumerate(prs.slides, 1):
    for sh in sl.shapes:
        if (sh.has_text_frame and sh.text_frame.text.strip().isdigit() and sh.top > ALT - Inches(0.6)
                and sh.left > LARG - Inches(1.3)):
            sh.text_frame.paragraphs[0].runs[0].text = str(i)
    tem = any(sh.has_text_frame and sh.text_frame.text.strip().isdigit() and sh.top > ALT - Inches(0.6)
              and sh.left > LARG - Inches(1.3) for sh in sl.shapes)
    if i > 1 and not tem:
        caixa(sl, LARG - Inches(1.0), ALT - Inches(0.46), Inches(0.5), Inches(0.3), str(i), 11, CLARO, alinha=PP_ALIGN.RIGHT)

prs.save(ARQ)
print("Salvo", ARQ, "| slides:", len(prs.slides))

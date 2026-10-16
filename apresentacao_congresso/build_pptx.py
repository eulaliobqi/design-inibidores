"""Monta a apresentacao de congresso (10 min, 15 slides) em PowerPoint 16:9.

Desenho: fundo branco, Segoe UI, paleta Okabe-Ito das figuras do manuscrito,
titulo de slide em frase afirmativa (assertion-evidence) e uma evidencia visual
por slide. Sem lista de topicos, sem icone decorativo, sem gradiente.

Uso (da raiz do projeto): python apresentacao_congresso/build_pptx.py
"""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

AQUI = Path(__file__).resolve().parent
FIG = AQUI / "figuras"
IMG = AQUI / "imagens"
SAIDA = AQUI / "Apresentacao_congresso_10min.pptx"

LARG, ALT = Inches(13.333), Inches(7.5)
MARGEM = Inches(0.78)

GRAF = RGBColor(0x1F, 0x23, 0x28)
AZUL = RGBColor(0x00, 0x72, 0xB2)
VERDE = RGBColor(0x2A, 0x9D, 0x8F)
LARA = RGBColor(0xD5, 0x5E, 0x00)
CINZA = RGBColor(0x5A, 0x5F, 0x66)
CLARO = RGBColor(0x8A, 0x90, 0x97)
FONTE = "Segoe UI"


def caixa(slide, x, y, cx, cy, texto, tam, cor=GRAF, negrito=False,
          alinha=PP_ALIGN.LEFT, italico=False, espaco=1.0, fonte=FONTE):
    tb = slide.shapes.add_textbox(x, y, cx, cy)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, linha in enumerate(texto.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = alinha
        p.line_spacing = espaco
        r = p.add_run()
        r.text = linha
        r.font.size = Pt(tam)
        r.font.bold = negrito
        r.font.italic = italico
        r.font.name = fonte
        r.font.color.rgb = cor
    return tb


def regua(slide, x, y, cx, cor=AZUL, espessura=Pt(3)):
    from pptx.enum.shapes import MSO_SHAPE
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, cx, espessura)
    s.fill.solid()
    s.fill.fore_color.rgb = cor
    s.line.fill.background()
    s.shadow.inherit = False
    return s


def nova(prs, numero=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    fundo = slide.background.fill
    fundo.solid()
    fundo.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    if numero is not None:
        caixa(slide, LARG - Inches(1.0), ALT - Inches(0.46), Inches(0.5), Inches(0.3),
              str(numero), 11, CLARO, alinha=PP_ALIGN.RIGHT)
    return slide


TOPO_FIG = Inches(2.05)


def titulo_slide(slide, frase, apoio=None):
    """Titulo em frase afirmativa; apoio e uma linha de contexto opcional.

    Define TOPO_FIG: onde a area util da figura comeca neste slide."""
    global TOPO_FIG
    caixa(slide, MARGEM, Inches(0.52), LARG - 2 * MARGEM, Inches(1.0),
          frase, 25, GRAF, negrito=True, espaco=1.08)
    regua(slide, MARGEM, Inches(0.40), Inches(0.9))
    if apoio:
        caixa(slide, MARGEM, Inches(1.52), LARG - 2 * MARGEM, Inches(0.6),
              apoio, 15, CINZA, espaco=1.18)
        TOPO_FIG = Inches(2.42) if "\n" in apoio or len(apoio) > 90 else Inches(2.20)
    else:
        TOPO_FIG = Inches(1.95)


LONGOS = []


def rodape(slide, texto):
    if len(texto) > 150:            # acima disso o rodape quebra em duas linhas
        LONGOS.append((len(texto), texto[:70]))
    caixa(slide, MARGEM, ALT - Inches(0.52), LARG - 2 * MARGEM - Inches(0.7), Inches(0.36),
          texto, 10.5, CLARO, espaco=1.12)


RODAPE_TOPO = Inches(6.90)


def figura(slide, arquivo, topo, largura_max=None):
    """Encaixa a figura entre `topo` e o rodape, centralizada e sem distorcer."""
    from PIL import Image
    p = FIG / arquivo if (FIG / arquivo).exists() else IMG / arquivo
    w, h = Image.open(p).size
    largura_max = largura_max or (LARG - 2 * MARGEM)
    altura_max = RODAPE_TOPO - topo
    esc = min(largura_max / w, altura_max / h)
    cx, cy = int(w * esc), int(h * esc)
    # centraliza tambem na vertical quando sobra altura
    y = topo + int((altura_max - cy) / 2)
    slide.shapes.add_picture(str(p), int((LARG - cx) / 2), y, cx, cy)


def notas(slide, texto):
    slide.notes_slide.notes_text_frame.text = texto


# ---------------------------------------------------------------------------
prs = Presentation()
prs.slide_width, prs.slide_height = LARG, ALT

# ------------------------------------------------------------- 1. capa ----
s = nova(prs)
regua(s, MARGEM, Inches(1.62), Inches(2.2), AZUL, Pt(4))
caixa(s, MARGEM, Inches(2.02), Inches(10.6), Inches(2.4),
      "Desenho Racional de Inibidores\nPeptídicos de Tripsinas contra\nLepidoptera por Inteligência\nArtificial Generativa",
      35, GRAF, negrito=True, espaco=1.12)
caixa(s, MARGEM, Inches(4.92), Inches(10.4), Inches(0.5),
      "Da tripsina digestiva ao peptídeo candidato ordenado: uma plataforma computacional",
      16.5, CINZA)
caixa(s, MARGEM, Inches(5.82), Inches(9.0), Inches(1.0),
      "Eulálio Santos\nUniversidade Federal de Viçosa  ·  eulalio.santos@ufv.br",
      14, GRAF, espaco=1.35)
caixa(s, LARG - MARGEM - Inches(2.6), Inches(6.28), Inches(2.6), Inches(0.3),
      "Outubro de 2026", 12, CLARO, alinha=PP_ALIGN.RIGHT)
notas(s, "15 s. Abertura: 'Vou contar como saímos de uma enzima do intestino de uma lagarta "
         "até peptídeos candidatos ordenados — e, principalmente, o que cada número desse caminho "
         "sustenta e o que não sustenta.'")

# -------------------------------------------------- 2. o problema ---------
s = nova(prs, 2)
titulo_slide(s, "A soja brasileira bateu recorde — e a lagarta come a folha antes da colheita")
from PIL import Image as _I
_w, _h = _I.open(IMG / "desfolha.png").size
_alt = Inches(3.9)
_lar = int(_alt * _w / _h)
s.shapes.add_picture(str(IMG / "desfolha.png"), MARGEM, Inches(2.18), _lar, _alt)
_x = MARGEM + _lar + Inches(0.6)
_cx2 = LARG - _x - MARGEM
caixa(s, _x, Inches(2.22), _cx2, Inches(0.8), "180,4 milhões de toneladas", 27, AZUL, negrito=True)
caixa(s, _x, Inches(2.86), _cx2, Inches(0.6),
      "de soja na safra 2025/26, a maior da série da Conab", 14.5, CINZA)
caixa(s, _x, Inches(3.62), _cx2, Inches(0.8), "metade da soja do mundo", 27, AZUL, negrito=True)
caixa(s, _x, Inches(4.26), _cx2, Inches(0.6),
      "é produzida na América do Sul", 14.5, CINZA)
_tb = caixa(s, _x, Inches(5.06), _cx2, Inches(1.2),
            "A lagarta-da-soja ", 15, GRAF, espaco=1.28)
_p = _tb.text_frame.paragraphs[0]
for _t, _it in [("Anticarsia gemmatalis", True),
                (" está entre os principais desfolhadores das Américas, e a desfolha pode "
                 "ser total em alta infestação.", False)]:
    _r = _p.add_run()
    _r.text = _t
    _r.font.size = Pt(15)
    _r.font.italic = _it
    _r.font.name = FONTE
    _r.font.color.rgb = GRAF
rodape(s, "Conab, set. 2026  ·  Carpane et al. 2022, PLOS ONE 17:e0271084  ·  imagem ilustrativa gerada por IA")
notas(s, "40 s. Gancho. Dizer os dois números grandes e deixar a imagem falar. "
         "Não afirmar perda percentual: a literatura consultada não dá um número único de perda para A. gemmatalis.")

# ----------------------------------------- 3. o controle atual ------------
s = nova(prs, 3)
titulo_slide(s, "O controle depende de Bt e de inseticida, e os dois exigem manejo de resistência",
             "É aí que entra uma terceira via: atacar a digestão da própria lagarta.")
col = (LARG - 2 * MARGEM - Inches(1.0)) / 3
for i, (tit, txt, cor) in enumerate([
    ("Inseticida químico",
     "Afeta organismos não alvo, pode\nprovocar surtos de pragas e\nseleciona populações resistentes.", LARA),
    ("Soja Bt (Cry1Ac)",
     "Evento de alta dose para a\nlagarta-da-soja, mas depende de\nrefúgio e de plano de manejo\npara durar.", LARA),
    ("Inibidores de protease",
     "Inibidores vegetais reduzem a\nsobrevivência larval, mas a lagarta\nse adapta induzindo proteases\ninsensíveis.", AZUL),
]):
    x = MARGEM + i * (col + Inches(0.5))
    regua(s, x, Inches(3.05), col, cor, Pt(2.5))
    caixa(s, x, Inches(3.32), col, Inches(0.5), tit, 18, cor, negrito=True)
    caixa(s, x, Inches(3.96), col, Inches(2.2), txt, 14.5, GRAF, espaco=1.30)
rodape(s, "Ongaratto et al. 2021  ·  Jongsma et al. 1995  ·  sem relato de resistência de campo desta espécie "
          "ao Cry1Ac no Brasil em 2025–2026")
notas(s, "35 s. Transição para o alvo. Ser explícito: não estou dizendo que a lagarta já é resistente ao Bt; "
         "estou dizendo que as duas táticas atuais carregam custo de manejo. Slide de corte se o tempo apertar.")

# --------------------------------------------------- 4. o alvo -----------
s = nova(prs, 4)
titulo_slide(s, "A digestão da lagarta depende de tripsinas num intestino fortemente alcalino")
figura(s, "fig_alvo.png", TOPO_FIG)
rodape(s, "Dow 1992  ·  Perona & Craik 1995  ·  Hedstrom 2002  ·  dados do painel: Tabelas 1 e S2 do manuscrito")
notas(s, "45 s. Explicar, para quem é leigo: a tripsina corta proteína sempre depois de uma lisina ou de uma arginina, "
         "porque o bolso S1 tem um aspartato negativo no fundo que atrai essas duas cadeias laterais positivas. "
         "Guardar esse fato: ele é o nó da história.")

# ------------------------------------------------ 5. a escala ------------
s = nova(prs, 5)
titulo_slide(s, "Inibidores naturais funcionam, mas são proteínas grandes — e os peptídeos curtos são fracos")
figura(s, "fig_escala.png", TOPO_FIG)
rodape(s, "Almeida Barbosa et al. 2021 (GORE1, GORE2)  ·  Mariano et al. 2026 (peptídeos da pró-região)  ·  "
          "faixas de comprimento: Seção 4.2 do manuscrito")
notas(s, "35 s. O ponto: existe um vão entre a proteína grande que funciona e o peptídeo curto que é barato de produzir. "
         "Potência milimolar é fraca — mil vezes mais fraca que um inibidor micromolar.")

# ---------------------------------------------- 6. o paradoxo ------------
s = nova(prs, 6)
titulo_slide(s, "E aqui está o nó: o resíduo que ancora é o resíduo que a enzima corta")
figura(s, "fig_paradoxo.png", TOPO_FIG)
rodape(s, "Laskowski & Kato 1980  ·  Laskowski & Qasim 2000  ·  Perona & Craik 1995")
notas(s, "45 s. CLÍMAX DA TENSÃO — falar devagar. Resistir à digestão e ancorar no alvo puxam em sentidos opostos. "
         "Um peptídeo clivado no intestino não chega ao alvo. Pausa aqui antes de virar para a solução.")

# -------------------------------------------- 7. a plataforma ------------
s = nova(prs, 7)
titulo_slide(s, "Montamos uma plataforma de seis etapas, do receptor ao candidato ordenado",
             "A não clivabilidade entra como exigência binária, não como critério desejável.")
figura(s, "fig_plataforma.png", TOPO_FIG)
rodape(s, "Figura 1 do manuscrito  ·  nenhum ensaio enzimático e nenhuma contrasseleção contra proteases "
          "não alvo foram feitos")
notas(s, "55 s. A VIRADA e a mensagem central da palestra. Uma linha por etapa. "
         "Frisar: a mesma plataforma roda para as oito espécies, e cada etapa é calibrada antes de ser usada.")

# ------------------------------------------- 8. a calibração -------------
s = nova(prs, 8)
titulo_slide(s, "Etapa 2: antes de confiar no escore, calibramos com inibidores reais e iscas embaralhadas",
             "O Boltz-2 é uma rede neural que prediz o complexo e devolve uma confiança — não uma afinidade.")
figura(s, "fig_calibracao.png", TOPO_FIG)
rodape(s, "Figura 2 e Tabela 2 do manuscrito  ·  o ApTI entrou na calibração sem isca pareada  ·  "
          "duas execuções com entrada idêntica correlacionam só ρ = 0,57")
notas(s, "45 s. O resultado que importa: o escore separa dentro do par, mas não como limiar absoluto — "
         "uma isca embaralhada do SFTI-1 chegou a 0,944. Por isso, adiante, cada candidato é comparado com "
         "controles pareados contra o mesmo receptor, e nunca lido como 'confiança 0,9 = liga'.")

# ------------------------------------------------- 9. o funil ------------
s = nova(prs, 9)
titulo_slide(s, "Etapas 3 e 4: de 22.066 sequências desenhadas sobram 527",
             "RFdiffusion desenhou 880 esqueletos; o ProteinMPNN escreveu as sequências; "
             "o critério duro removeu todas que têm um resíduo que o intestino corta.")
figura(s, "fig_funil.png", TOPO_FIG)
rodape(s, "Figura 3 do manuscrito  ·  critério duro: nenhum dos nove resíduos que as proteases do intestino "
          "assumem em P1 (K, R, F, Y, W, L, M, A, V)")
notas(s, "45 s. Dizer o que cada ferramenta faz em uma linha. O número que fica: 97,6% do que a IA desenhou "
         "é descartado por uma regra biológica, não por um escore.")

# -------------------------------------------- 10. o preço do filtro ------
s = nova(prs, 10)
titulo_slide(s, "O preço do filtro: sobram peptídeos curtos e quase metade glicina")
figura(s, "fig_composicao.png", TOPO_FIG)
rodape(s, "Figura S3 do manuscrito  ·  dos 4.060 resíduos da frente L, só nove são lisina ou arginina (0,22%), "
          "cada um deles seguido de prolina")
notas(s, "35 s. É a consequência direta do paradoxo do slide 6: proibir os resíduos que a enzima corta é "
         "proibir justamente os que ancoram melhor. Slide de corte se o tempo apertar.")

# ------------------------------------------ 11. a pose inicial decide ----
s = nova(prs, 11)
titulo_slide(s, "Etapa 5: as simulações mostraram que quem fica no bolso é quem já começou dentro dele",
             "48 complexos, 10 ns cada, uma execução por pH.")
figura(s, "fig_pose_decide.png", TOPO_FIG)
rodape(s, "Figura 5 do manuscrito  ·  teste exato de Fisher, P = 7 × 10⁻⁵  ·  a âncora é definida post hoc")
notas(s, "45 s. Autocrítica honesta da plataforma: a simulação de 10 ns verifica se uma pose já proposta "
         "sobrevive; ela não dá tempo para o peptídeo achar o bolso. Por isso a ocupância deixou de ser critério "
         "e passou a ser descrição.")

# ---------------------------------------------------- 12. pH ------------
s = nova(prs, 12)
titulo_slide(s, "Entre pH 8,2 e pH 10,0 não há diferença — e só dois peptídeos repetem o resultado",
             "O intestino da lagarta é alcalino, mas o ensaio enzimático roda em pH 8,2: o candidato "
             "precisa funcionar na faixa inteira.")
figura(s, "fig_ph.png", TOPO_FIG)
rodape(s, "Figura 7 do manuscrito  ·  Wilcoxon pareado, n = 48  ·  repetir no mesmo pH já mudou a ocupância "
          "em mais de 0,5 em 3 de 16 pares")
notas(s, "40 s. Ponto delicado: a variação entre duas execuções do mesmo pH é do mesmo tamanho que a variação "
         "entre os dois pH. Então nenhuma mudança pode ser lida como efeito do pH. NGGRPDAP e GGKPGEP são os "
         "únicos que repetem.")

# -------------------------------------------- 13. os candidatos ---------
s = nova(prs, 13)
titulo_slide(s, "Etapa 6: quatro peptídeos sobreviveram a todos os filtros disponíveis")
figura(s, "fig_poses.png", TOPO_FIG)
rodape(s, "Figura 8 do manuscrito  ·  poses iniciais preditas pelo Boltz-2  ·  nenhum deles carrega alegação "
          "de seletividade ou de atividade inibitória")
notas(s, "45 s. Apresentar os quatro. Destacar os dois de A. gemmatalis: NGGRPDAP, linear, com arginina em S1 "
         "seguida de prolina, e GGKPGEP, cíclico, com lisina e a menor distância final ao Asp189 (2,71 Å).")

# --------------------------------------------- 14. a prolina ------------
s = nova(prs, 14)
titulo_slide(s, "A prolina depois da lisina ou da arginina protege o sítio nas simulações",
             "Um controle embaralhado com a mesma composição do NGGRPDAP dá a linha de base interna.")
figura(s, "fig_clivagem.png", TOPO_FIG)
rodape(s, "Figura S11 do manuscrito  ·  os dois ocupam S1 igualmente (1,00) e diferem só no resíduo seguinte  "
          "·  não é uma taxa de clivagem")
notas(s, "40 s. Fechamento do arco do slide 6: o paradoxo tem uma saída, e ela é geométrica. "
         "Ressalva obrigatória: a clivagem antes de prolina é mais lenta, não ausente, e o critério aqui é "
         "geométrico, não cinético. A aminopeptidase N continua sendo um furo no peptídeo linear.")

# ------------------------------------------------- 15. fecho ------------
s = nova(prs, 15)
titulo_slide(s, "A plataforma entrega candidatos ordenados; quem decide se eles funcionam é o ensaio enzimático")
_w, _h = _I.open(IMG / "ensaio.png").size
_alt = Inches(3.5)
_lar = int(_alt * _w / _h)
s.shapes.add_picture(str(IMG / "ensaio.png"), LARG - MARGEM - _lar, Inches(2.35), _lar, _alt)
_cx = LARG - 2 * MARGEM - _lar - Inches(0.55)
caixa(s, MARGEM, Inches(2.25), _cx, Inches(0.4), "O que a plataforma entrega", 17, AZUL, negrito=True)
caixa(s, MARGEM, Inches(2.72), _cx, Inches(1.5),
      "Um caminho reprodutível de oito tripsinas até peptídeos\nordenados, com a não clivabilidade como exigência\nexplícita e cada escore calibrado antes de ser usado.",
      14.5, GRAF, espaco=1.28)
caixa(s, MARGEM, Inches(4.28), _cx, Inches(0.4), "O que ela ainda não diz", 17, LARA, negrito=True)
caixa(s, MARGEM, Inches(4.75), _cx, Inches(1.2),
      "Tudo é computacional. Não houve ensaio de inibição, de\nestabilidade nem contrasseleção contra proteases não alvo.",
      14.5, GRAF, espaco=1.28)
caixa(s, MARGEM, Inches(5.58), _cx, Inches(0.4), "O próximo passo", 17, VERDE, negrito=True)
_tb = caixa(s, MARGEM, Inches(6.02), _cx, Inches(0.8),
            "Inibição e estabilidade em extrato de intestino de ", 14.5, GRAF, espaco=1.28)
_p = _tb.text_frame.paragraphs[0]
for _t, _it in [("A. gemmatalis", True), (", com o GORE2 como controle positivo.", False)]:
    _r = _p.add_run()
    _r.text = _t
    _r.font.size = Pt(14.5)
    _r.font.italic = _it
    _r.font.name = FONTE
    _r.font.color.rgb = GRAF
rodape(s, "github.com/eulaliobqi/design-inibidores  ·  eulalio.santos@ufv.br  ·  imagem ilustrativa gerada por IA")
notas(s, "40 s. Take-home: a plataforma existe, roda e é transferível para outro alvo. "
         "Terminar na frase: 'um candidato só interessa se o Ki medido nas mesmas condições ficar abaixo da faixa "
         "milimolar do GORE2, ou se ele persistir mais tempo no extrato de intestino.' Agradecer e abrir para perguntas.")

for n, t in LONGOS:
    print(f"  [rodape longo: {n} car.] {t}...")
prs.save(SAIDA)
print(f"Salvo: {SAIDA}")
print(f"Slides: {len(prs.slides.__iter__.__self__._sldIdLst)}")

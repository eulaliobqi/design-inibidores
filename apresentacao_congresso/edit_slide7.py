"""Formata o slide 7 depois que o autor trocou as figuras (resumo grafico, curva de sobrevivencia, pupas).
Preserva as figuras do autor: so reposiciona/redimensiona (proporcao mantida), ajusta titulo, legendas e rodape.
"""
import shutil
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

AQUI = Path(__file__).resolve().parent
ARQ = AQUI / "Apresentacao_congresso_v2.pptx"
shutil.copy2(ARQ, AQUI / "Apresentacao_congresso_v2_antes_slide7.pptx")

CLARO = RGBColor(0x8A, 0x90, 0x97)
CINZA = RGBColor(0x5A, 0x5F, 0x66)
FONTE = "Segoe UI"
LARG = Inches(13.333)


def legenda(slide, x, y, cx, texto, tam=11.5, cor=CINZA, alinha=PP_ALIGN.CENTER):
    tb = slide.shapes.add_textbox(x, y, cx, Inches(0.28))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = alinha
    r = p.add_run()
    r.text = texto
    r.font.size = Pt(tam)
    r.font.name = FONTE
    r.font.color.rgb = cor
    return tb


prs = Presentation(ARQ)
s = prs.slides[6]
pics = [sh for sh in s.shapes if sh.shape_type == 13]
assert len(pics) == 3, len(pics)
for p in pics:
    p.lock_aspect_ratio = True
asp = lambda p: p.width / p.height
resumo = min(pics, key=lambda p: abs(asp(p) - 1.15))
surv = [p for p in pics if p is not resumo and asp(p) > 1.55][0]
pupas = [p for p in pics if p not in (resumo, surv)][0]
print("resumo", round(asp(resumo), 2), "curva", round(asp(surv), 2), "pupas", round(asp(pupas), 2))


def coloca(p, x, y, altura):
    p.height = int(altura)
    p.width = int(altura * asp_orig[id(p)])
    p.left, p.top = int(x), int(y)


asp_orig = {id(p): asp(p) for p in pics}
# coluna esquerda: resumo grafico
coloca(resumo, Inches(0.78), Inches(1.85), Inches(4.45))
legenda(s, Inches(0.78), Inches(1.85) + Inches(4.45) + Inches(0.08), resumo.width,
        "Resumo do experimento: GORE3 sobre as plantas e as lagartas")
# coluna direita: curva e pupas
xr = Inches(6.55)
cw = LARG - Inches(0.78) - xr
coloca(surv, xr + int((cw - Inches(2.6) * asp_orig[id(surv)]) / 2), Inches(1.8), Inches(2.6))
legenda(s, xr, Inches(1.8) + Inches(2.6) + Inches(0.04), cw, "Sobrevivência das lagartas ao longo do tempo")
coloca(pupas, xr + int((cw - Inches(1.65) * asp_orig[id(pupas)]) / 2), Inches(4.93), Inches(1.65))
legenda(s, xr, Inches(4.93) + Inches(1.65) + Inches(0.04), cw, "Pupas (barra: 2 mm)")

# titulo coerente com as novas figuras
for sh in s.shapes:
    if sh.has_text_frame and sh.text_frame.text.startswith("No nosso grupo"):
        sh.text_frame.paragraphs[0].runs[0].text = (
            "No nosso grupo, o peptídeo GORE3 reduz a sobrevivência da lagarta mais que o inibidor da soja (SKTI)")
        for r in sh.text_frame.paragraphs[0].runs[1:]:
            r.text = ""

# rodape
tb = s.shapes.add_textbox(Inches(0.78), Inches(7.5) - Inches(0.55), LARG - Inches(0.78) * 2 - Inches(0.7), Inches(0.4))
tf = tb.text_frame
tf.word_wrap = True
tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
r = tf.paragraphs[0].add_run()
r.text = ("Figuras: dados do grupo (Laboratório de Enzimologia, Bioquímica de Proteínas e Peptídeos I e II, UFV/BIOAGRO/DBB)  ·  "
          "família GORE: Almeida Barros et al. 2021, Pest Manag. Sci. 77:1714  ·  Júnior et al. 2026, Arch. Insect Biochem. Physiol. 123:e70223")
r.font.size = Pt(10)
r.font.name = FONTE
r.font.color.rgb = CLARO
tf = s.notes_slide.notes_text_frame
if tf is not None:
    tf.text = ("45 s. Resultados do grupo com o peptídeo GORE3: esquema do experimento (plantas tratadas, lagartas alimentadas, "
               "inibição digestiva, estresse oxidativo, dano celular e queda da sobrevivência), curva de sobrevivência (GORE3 abaixo "
               "do controle, da benzamidina e do SKTI) e pupas. Confirmar espécie, concentração e referência da publicação desta figura.")
prs.save(ARQ)
print("ok")

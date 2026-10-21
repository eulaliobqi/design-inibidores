"""2a rodada de ajustes, em cima do arquivo editado pelo autor:
- slide 2 (ODS): so o icone e UMA frase
- novo slide 10 (apos o fluxograma): animacao RFdiffusion -> ProteinMPNN -> Boltz-2 -> MD
- renumera os slides
Uso: python apresentacao_congresso/edit_sia2.py   (com o PowerPoint fechado)
"""
import shutil
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

AQUI = Path(__file__).resolve().parent
ARQ = AQUI / "Apresentacao_congresso_v2.pptx"
shutil.copy2(ARQ, AQUI / "Apresentacao_congresso_v2_antes_rodada2.pptx")

LARG, ALT = Inches(13.333), Inches(7.5)
MARGEM = Inches(0.78)
GRAF = RGBColor(0x1F, 0x23, 0x28)
AZUL = RGBColor(0x00, 0x72, 0xB2)
AMAR = RGBColor(0xA8, 0x70, 0x00)
CLARO = RGBColor(0x8A, 0x90, 0x97)
FONTE = "Segoe UI"


def caixa(slide, x, y, cx, cy, texto, tam, cor=GRAF, negrito=False, alinha=PP_ALIGN.LEFT, espaco=1.0):
    tb = slide.shapes.add_textbox(x, y, cx, cy)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, linha in enumerate(texto.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = alinha
        p.line_spacing = espaco
        r = p.add_run()
        r.text = linha
        r.font.size = Pt(tam)
        r.font.bold = negrito
        r.font.name = FONTE
        r.font.color.rgb = cor
    return tb


def regua(slide, x, y, cx, cor=AZUL, espessura=Pt(3)):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, cx, espessura)
    s.fill.solid()
    s.fill.fore_color.rgb = cor
    s.line.fill.background()
    s.shadow.inherit = False
    return s


def notas(slide, texto):
    tf = slide.notes_slide.notes_text_frame
    if tf is not None:
        tf.text = texto


prs = Presentation(ARQ)

# ============================================================ slide 2 (ODS) ====
s2 = prs.slides[1]
for sh in list(s2.shapes):
    manter = (sh.shape_type == 13
              or (sh.has_text_frame and (sh.text_frame.text.strip().isdigit() or sh.text_frame.text.startswith("ONU Brasil")))
              or (not sh.has_text_frame and sh.shape_type == 1 and sh.top < Inches(0.6)))
    if not manter:
        sh._element.getparent().remove(sh._element)
for sh in s2.shapes:
    if sh.shape_type == 13:                                   # icone da ODS, maior e centrado na vertical
        sh.left, sh.top, sh.width, sh.height = MARGEM, Inches(1.45), Inches(4.8), Inches(4.8)
    elif sh.has_text_frame and sh.text_frame.text.startswith("ONU Brasil"):
        sh.text_frame.paragraphs[0].runs[0].text = (
            "ONU Brasil, ODS 2 (brasil.un.org/pt-br/sdgs/2)  ·  ícone oficial: sdgs.un.org  ·  "
            "meta 2.4: sistemas sustentáveis de produção de alimentos e práticas agrícolas resilientes")
x = MARGEM + Inches(5.6)
frase = caixa(s2, x, Inches(2.35), LARG - x - MARGEM, Inches(3.0),
              "Controlar lagartas-praga com peptídeos desenhados por inteligência artificial, "
              "para proteger a produtividade agrícola de forma sustentável.", 30, GRAF, negrito=True, espaco=1.18)
regua(s2, x, Inches(2.1), Inches(0.9), AMAR, Pt(4))

# ================================ novo slide 10: como as ferramentas atuam ====
novo = prs.slides.add_slide(prs.slide_layouts[6])
novo.background.fill.solid()
novo.background.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
caixa(novo, MARGEM, Inches(0.52), LARG - 2 * MARGEM, Inches(1.0),
      "Como as ferramentas atuam, passo a passo: esqueleto, sequência, complexo e dinâmica", 25, GRAF, True, espaco=1.08)
regua(novo, MARGEM, Inches(0.40), Inches(0.9))
vw = Inches(11.2)
vh = int(vw * 490 / 1120)
vx = int((LARG - vw) / 2)
novo.shapes.add_movie(str(AQUI / "imagens" / "como_funciona_RFdiff_MPNN_Boltz_MD.mp4"), vx, Inches(1.65), vw, vh,
                      poster_frame_image=str(AQUI / "imagens" / "como_funciona_poster.png"), mime_type="video/mp4")
caixa(novo, MARGEM, ALT - Inches(0.55), LARG - 2 * MARGEM - Inches(0.7), Inches(0.4),
      "RFdiffusion: Watson et al. 2023  ·  ProteinMPNN: Dauparas et al. 2022  ·  Boltz-2: Passaro et al. 2025  ·  "
      "GROMACS: Abraham et al. 2015  ·  exemplo real (GGKPGEP): passos de denoising, sequência e quadros de MD são dados do projeto",
      10, CLARO, espaco=1.1)
notas(novo, "45 s. Um exemplo real, o GGKPGEP de A. gemmatalis. 1) O RFdiffusion parte de ruído e, em 50 passos, desenha o "
            "esqueleto do peptídeo (um anel de 7 resíduos) no bolso da tripsina. 2) O ProteinMPNN escreve a sequência que dobra "
            "nesse esqueleto: G G K P G E P. 3) O Boltz-2 prediz o complexo peptídeo + tripsina, com o encaixe da lisina no bolso S1 "
            "(confiança 0,92, que só vale para comparar com controles). 4) A dinâmica molecular de 10 ns testa se essa pose se mantém. "
            "A pose do Boltz-2 pode diferir da posição do esqueleto do RFdiffusion: cada etapa é independente.")
lst = prs.slides._sldIdLst
ids = list(lst)
no = ids[-1]
lst.remove(no)
lst.insert(9, no)                   # logo apos o slide 9 (fluxograma)

# ================================================================ renumera ====
for i, sl in enumerate(prs.slides, 1):
    tem = False
    for sh in sl.shapes:
        if (sh.has_text_frame and sh.text_frame.text.strip().isdigit() and sh.top > ALT - Inches(0.6)
                and sh.left > LARG - Inches(1.3)):
            sh.text_frame.paragraphs[0].runs[0].text = str(i)
            tem = True
    if i > 1 and not tem and sl.shapes:
        caixa(sl, LARG - Inches(1.0), ALT - Inches(0.46), Inches(0.5), Inches(0.3), str(i), 11, CLARO, alinha=PP_ALIGN.RIGHT)

prs.save(ARQ)
print("Salvo", ARQ, "| slides:", len(prs.slides))

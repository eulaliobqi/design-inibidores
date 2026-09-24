"""Acrescenta slides de atualizacao (estado atual, 2026-09-23) ao final da
apresentacao existente, SEM regenerar do zero — preserva qualquer edicao
manual feita direto no .pptx desde o ultimo commit (ex.: ajustes de 2026-07-21
nunca commitados). scripts/gerar_apresentacao_pptx.py regenera as 22 slides
originais a partir de scripts/canva_slides_data.json (que nao tem esse
conteudo novo) — rodar aquele script aqui apagaria a edicao manual pendente
sem ganhar nada em troca.

Uso: python -m scripts.append_update_slides
"""
from pathlib import Path

from pptx import Presentation

from scripts.gerar_apresentacao_pptx import add_content_slide

PPTX_PATH = Path("apresentacao_design_inibidores.pptx")

NEW_SLIDES = [
    (
        "Calibração da Escada de Scoring (B0.5) — Fechada",
        [
            "6 inibidores reais (BPTI, SFTI-1, SKTI, Bowman-Birk, EcTI, ApTI) + 10 decoys, "
            "2 receptores (tripsina bovina real + S. frugiperda)",
            "Boltz-2 (confidence/pLDDT): separa real de decoy em 10/10 pares — único método validado",
            "RMSD de MD curta (2ns, complexo inteiro): 5/10 — equivalente a acaso",
            "MM-PBSA (GB, trajetória única): 4/10 — PIOR que acaso (decoy vence por até +144,8 kcal/mol)",
            "Decisão do usuário: loop de contrasseleção (B2.7) usa apenas Boltz-2 como scorer",
        ],
    ),
    (
        "B1.4 — Mapeamento Real dos Subsítios S1-S4/S1'-S3'",
        [
            "Substitui heurística quebrada (offset sistemático +15/16 resíduos em 8/9 espécies)",
            "P1 de BPTI/SFTI-1 achado por geometria real em complexos cristalográficos reais (2PTC, 1SFI)",
            "Transferência para os 9 receptores via foldseek TMalign — equivalência estrutural real",
            "Resultado: 18/18 espécie×template aprovados (TMscore ~0,95, RMSD ~1,3Å)",
            "Validação cruzada: 100% de concordância com método independente anterior (offset de sequência)",
        ],
    ),
    (
        "Pivô: Geração por Difusão contra Lepidoptera (em andamento)",
        [
            "Pausa temporária da linha de contrasseleção (painel negativo/seletividade) — foco em "
            "potência ampla contra o painel de 7 pragas-alvo",
            "Trilha A: macrociclo via RFdiffusion cíclico, ancorado nos hotspots reais de B1.4",
            "Piloto técnico validado ponta-a-ponta (fechamento N-C real ~1,3Å, ProteinMPNN funcionando)",
            "Campanha real em andamento: 7 espécies × 5 comprimentos (8-16aa) × 10 designs = 350 backbones",
            "Próximo passo: triagem dos candidatos gerados via Boltz-2 (único scorer validado)",
        ],
    ),
]


def main():
    prs = Presentation(str(PPTX_PATH))
    n_before = len(prs.slides)
    for title, bullets in NEW_SLIDES:
        add_content_slide(prs, title, bullets)
    prs.save(str(PPTX_PATH))
    print(f"Slides antes: {n_before} | depois: {len(prs.slides)} | salvo em {PPTX_PATH.resolve()}")


if __name__ == "__main__":
    main()
